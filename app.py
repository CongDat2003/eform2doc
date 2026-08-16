"""
app.py — Web UI cho eform2doc.

Chạy local:  streamlit run app.py
Deploy free: Streamlit Community Cloud / Hugging Face Spaces
"""

import io
import json
import time

import streamlit as st

from eform_parser import (
    parse_eform, all_placeholders, GRID_TYPES, MULTI_CHECKBOX_TYPES,
)
from doc_tools import scan_placeholders, generate_template
from auto_fill import find_slots, match_fields, apply_suggestions, unmatched_fields

st.set_page_config(page_title="eform2doc", page_icon="📄", layout="wide")

st.title("📄 eform2doc")
st.caption("Sinh và kiểm tra DOC template từ eform JSON (Form.io)")


# ------------------------------------------------------------------ helpers

def fields_to_rows(fields):
    grid_keys = {f.key for f in fields if f.type in GRID_TYPES}
    rows = []
    for f in fields:
        if f.parent in grid_keys:
            rows.append({
                "Key": f"    └ {f.key}",
                "Nhãn": f.label,
                "Loại": f.type,
                "Bắt buộc": "✓" if f.required else "",
                "Placeholder": f"cột của ${{{f.parent}#table}}",
            })
        elif f.type in MULTI_CHECKBOX_TYPES:
            rows.append({
                "Key": f.key, "Nhãn": f.label, "Loại": f.type,
                "Bắt buộc": "✓" if f.required else "",
                "Placeholder": f"{len(f.options)} option ↓",
            })
            for val, lbl, ph in f.option_placeholders():
                rows.append({
                    "Key": f"    · {val}", "Nhãn": lbl, "Loại": "",
                    "Bắt buộc": "", "Placeholder": ph,
                })
        else:
            rows.append({
                "Key": f.key, "Nhãn": f.label, "Loại": f.type,
                "Bắt buộc": "✓" if f.required else "",
                "Placeholder": f.placeholder,
            })
    return rows


def build_markdown(fields):
    grid_keys = {f.key for f in fields if f.type in GRID_TYPES}
    lines = ["# Bảng mapping eform ↔ DOC template", "",
             "| Key | Nhãn | Loại | Bắt buộc | Placeholder |", "|---|---|---|---|---|"]
    for f in fields:
        if f.parent in grid_keys:
            key_txt, ph = f"&nbsp;&nbsp;└ `{f.key}`", f"cột của `${{{f.parent}#table}}`"
        else:
            key_txt, ph = f"`{f.key}`", f"`{f.placeholder}`"
        label = (f.label or "").replace("|", "\\|")
        lines.append(f"| {key_txt} | {label} | {f.type} | {'✓' if f.required else ''} | {ph} |")

    multi = [f for f in fields if f.type in MULTI_CHECKBOX_TYPES]
    if multi:
        lines += ["", "## Chi tiết nhóm checkbox", ""]
        for f in multi:
            lines += [f"### `{f.key}` — {f.label}", "",
                      "| Value | Nhãn | Placeholder |", "|---|---|---|"]
            for val, lbl, ph in f.option_placeholders():
                lines.append(f"| `{val}` | {lbl.replace('|', chr(92) + '|')} | `{ph}` |")
            lines.append("")
    return "\n".join(lines) + "\n"


# ------------------------------------------------------------------ input

uploaded = st.file_uploader("Tải file eform JSON", type=["json"])

with st.expander("Hoặc dán JSON trực tiếp"):
    pasted = st.text_area("JSON", height=180, label_visibility="collapsed")

raw = None
source_name = "form"
if uploaded is not None:
    raw = json.load(uploaded)
    source_name = uploaded.name.rsplit(".", 1)[0]
elif pasted.strip():
    try:
        raw = json.loads(pasted)
    except json.JSONDecodeError as e:
        st.error(f"JSON không hợp lệ: {e}")

if raw is None:
    st.info("Tải lên hoặc dán eform JSON để bắt đầu.")
    st.stop()


# ------------------------------------------------------------------ parse

t0 = time.perf_counter()
fields = parse_eform(raw)
parse_ms = (time.perf_counter() - t0) * 1000

placeholders = all_placeholders(fields)
grids = [f for f in fields if f.type in GRID_TYPES]

c1, c2, c3, c4 = st.columns(4)
c1.metric("Field", len(fields))
c2.metric("Placeholder", len(placeholders))
c3.metric("Bảng (datagrid)", len(grids))
c4.metric("Thời gian đọc", f"{parse_ms:.0f} ms")

if grids:
    st.success("Bảng phát hiện được: " + " · ".join(
        f"`${{{g.key}#table}}`" for g in grids))

tab0, tab1, tab2, tab3 = st.tabs([
    "⚡ Cấu hình DOC gốc",
    "📋 Danh sách key",
    "📄 Sinh DOC mới",
    "🔍 Đối chiếu DOC",
])


# ------------------------------------------------------------------ tab 0

with tab0:
    st.write(
        "Tải lên **file DOC gốc** (mẫu giấy, chưa có placeholder). "
        "Tool tự dò các chỗ trống, ghép với field trong eform và chèn `${key}` "
        "vào đúng vị trí — **giữ nguyên toàn bộ layout gốc**."
    )

    doc_src = st.file_uploader("DOC gốc (.docx)", type=["docx"], key="autofill_src")

    min_score = st.slider(
        "Ngưỡng khớp tối thiểu", 0.30, 0.90, 0.45, 0.05,
        help="Thấp hơn = ghép được nhiều hơn nhưng dễ sai. Cao hơn = chắc chắn hơn nhưng bỏ sót.",
    )

    tbl_mode_label = st.radio(
        "Cách cấu hình bảng trong DOC",
        [
            "Xoá bảng, thay bằng ${TenBang#table}",
            "Giữ bảng, mỗi cột một ${TenBang.TenCot}",
        ],
        help="Chọn theo cách backend của bạn merge bảng.",
    )
    tbl_mode = "replace" if tbl_mode_label.startswith("Xoá") else "row"

    do_tidy = st.checkbox(
        "Tự căn chỉnh lại tài liệu sau khi chèn",
        value=True,
        help="Tiêu đề bảng in đậm căn giữa và lặp khi sang trang, ô bảng căn "
             "giữa theo chiều dọc, gộp dòng trống thừa, thống nhất giãn dòng.",
    )

    if doc_src is not None:
        import os
        import tempfile

        from docx import Document as _Doc

        work = tempfile.mkdtemp()
        src_path = os.path.join(work, "goc.docx")
        with open(src_path, "wb") as fh:
            fh.write(doc_src.getbuffer())

        t0 = time.perf_counter()
        slots = find_slots(_Doc(src_path))
        sugs = match_fields(slots, fields, threshold=min_score)
        match_ms = (time.perf_counter() - t0) * 1000

        matched = [s for s in sugs if s.field]
        skipped = [s for s in sugs if not s.field]

        a, b, c, d = st.columns(4)
        a.metric("Chỗ trống tìm thấy", len(slots))
        b.metric("Ghép được", len(matched))
        c.metric("Không khớp", len(skipped))
        d.metric("Thời gian dò", f"{match_ms:.0f} ms")

        low = [s for s in matched if s.score < 0.60]
        if low:
            st.warning(
                f"{len(low)} cặp có độ tin cậy thấp — bỏ tick nếu thấy sai."
            )

        st.markdown("#### Duyệt các cặp ghép")
        st.caption("Bỏ tick ở cột **Áp dụng** với cặp nào bạn thấy không đúng.")

        rows = []
        for i, s in enumerate(matched):
            rows.append({
                "Áp dụng": True,
                "Tin cậy": s.confidence,
                "Điểm": round(s.score, 2),
                "Vị trí trong DOC": s.slot.location,
                "Nhãn đọc được": s.slot.label[:60],
                "Placeholder": s.placeholder,
            })

        edited = st.data_editor(
            rows,
            use_container_width=True,
            height=min(420, 60 + 36 * max(len(rows), 1)),
            column_config={
                "Áp dụng": st.column_config.CheckboxColumn(width="small"),
                "Tin cậy": st.column_config.TextColumn(width="small"),
                "Điểm": st.column_config.NumberColumn(width="small", format="%.2f"),
            },
            disabled=["Tin cậy", "Điểm", "Vị trí trong DOC",
                      "Nhãn đọc được", "Placeholder"],
            hide_index=True,
            key="review_table",
        )

        col_l, col_r = st.columns(2)

        with col_l:
            if skipped:
                with st.expander(f"Chỗ trống không khớp field nào ({len(skipped)})"):
                    for s in skipped:
                        st.caption(f"**{s.slot.location}** — {s.slot.label[:70]}")

        with col_r:
            um = unmatched_fields(fields, sugs)
            if um:
                with st.expander(f"Field eform chưa có chỗ trong DOC ({len(um)})"):
                    for f, ph in um:
                        st.caption(f"`{ph}` — {f.label[:55]}")

        st.divider()

        if st.button("Chèn placeholder vào DOC gốc", type="primary"):
            accepted_matched = [bool(r["Áp dụng"]) for r in edited]
            accepted = []
            mi = 0
            for s in sugs:
                if s.field:
                    accepted.append(accepted_matched[mi])
                    mi += 1
                else:
                    accepted.append(False)

            t0 = time.perf_counter()
            out_path = os.path.join(work, "da_cau_hinh.docx")
            n = apply_suggestions(src_path, out_path, sugs, accepted,
                                  all_fields=fields, tidy=do_tidy,
                                  table_mode=tbl_mode)
            apply_ms = (time.perf_counter() - t0) * 1000

            after = scan_placeholders(out_path)
            st.success(
                f"Đã chèn {n} placeholder trong {apply_ms:.0f} ms — "
                f"kiểm tra lại thấy {len(after)} placeholder trong file."
            )

            need = set(all_placeholders(fields,
                                        with_columns=(tbl_mode == 'row')))
            still_missing = sorted(need - after)
            if still_missing:
                st.info(
                    f"Còn {len(still_missing)} placeholder chưa có trong DOC "
                    "(field ẩn hoặc DOC gốc không có mục tương ứng):"
                )
                st.code("\n".join(still_missing), language="text")

            with open(out_path, "rb") as fh:
                st.download_button(
                    "⬇ Tải DOC đã cấu hình",
                    fh.read(),
                    file_name=f"{source_name}_da_cau_hinh.docx",
                    mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                )
    else:
        st.info("Tải lên DOC gốc để bắt đầu dò chỗ trống.")


# ------------------------------------------------------------------ tab 1

with tab1:
    st.dataframe(fields_to_rows(fields), use_container_width=True, height=460)
    st.download_button(
        "⬇ Tải bảng mapping (.md)",
        build_markdown(fields),
        file_name=f"{source_name}_mapping.md",
        mime="text/markdown",
    )


# ------------------------------------------------------------------ tab 2

with tab2:
    title = st.text_input("Tiêu đề văn bản", value="BIỂU MẪU")
    subtitle = st.text_input("Phụ đề (không bắt buộc)", value="")

    if st.button("Sinh DOC template", type="primary"):
        t0 = time.perf_counter()
        buf = io.BytesIO()
        import tempfile, os
        tmp = os.path.join(tempfile.mkdtemp(), "out.docx")
        generate_template(fields, tmp, title=title, subtitle=subtitle)
        with open(tmp, "rb") as fh:
            buf.write(fh.read())
        gen_ms = (time.perf_counter() - t0) * 1000

        st.success(f"Xong trong {gen_ms:.0f} ms — {len(placeholders)} placeholder")
        st.download_button(
            "⬇ Tải template.docx",
            buf.getvalue(),
            file_name=f"{source_name}_template.docx",
            mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        )

    with st.expander("Danh sách placeholder sẽ có trong DOC"):
        st.code("\n".join(sorted(placeholders)), language="text")


# ------------------------------------------------------------------ tab 3

with tab3:
    st.write("Tải lên DOC template đang dùng để kiểm tra có lệch key với eform không.")
    doc_file = st.file_uploader("DOC template (.docx)", type=["docx"], key="check")

    if doc_file is not None:
        import tempfile, os
        tmp = os.path.join(tempfile.mkdtemp(), "check.docx")
        with open(tmp, "wb") as fh:
            fh.write(doc_file.read())

        t0 = time.perf_counter()
        found = scan_placeholders(tmp)
        scan_ms = (time.perf_counter() - t0) * 1000

        needed = set(placeholders)
        missing = sorted(needed - found)
        extra = sorted(found - needed)
        ok = needed & found

        m1, m2, m3, m4 = st.columns(4)
        m1.metric("eform cần", len(needed))
        m2.metric("DOC đang có", len(found))
        m3.metric("Khớp", len(ok))
        m4.metric("Thời gian", f"{scan_ms:.0f} ms")

        if not missing and not extra:
            st.success("Khớp hoàn toàn — DOC template sẵn sàng merge.")
        else:
            if missing:
                st.error(f"THIẾU {len(missing)} placeholder — eform có nhưng DOC chưa có:")
                st.code("\n".join(missing), language="text")
            if extra:
                st.warning(f"THỪA {len(extra)} placeholder — DOC có nhưng eform không sinh ra:")
                st.code("\n".join(extra), language="text")


st.divider()
st.caption(
    "Quy tắc: `textfield/select/datetime → ${Key}` · "
    "`checkbox → ${Key#checkbox}` · "
    "`selectboxes → ${Key.value#checkbox}` · "
    "`datagrid/editgrid → ${TenBang#table}`"
)
